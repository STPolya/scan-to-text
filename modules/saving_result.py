from modules.logging_config import log_error
import pytesseract
from docx import Document
import xlsxwriter

def save_to_txt(text, file_path_to_save):
    try:
        with open(file_path_to_save, 'w', encoding='utf-8') as file:
            file.write(text)
        return True

    except Exception as error:
        log_error('saving_to_txt', file_path_to_save, error)
        return f'Не удалось сохранить файл {error}'


def save_to_pdf(image, file_path_to_save):
    try:
        pdf = pytesseract.image_to_pdf_or_hocr(image, extension='pdf')
        with open(file_path_to_save, 'wb') as f:
            f.write(pdf)
        return True

    except Exception as error:
        log_error('saving_to_pdf', file_path_to_save, error)
        return  f'Не удалось сохранить файл {error}'


def save_to_docx(text, file_path_to_save):
    try:
        doc = Document()
        doc.add_paragraph(text)
        doc.save(file_path_to_save)
        return True

    except Exception as error:
        log_error('saving_to_docx', file_path_to_save, error)
        return f'Не удалось сохранить файл {error}'


def save_to_xlsx(table, file_path_to_save):
    try:
        workbook = xlsxwriter.Workbook(file_path_to_save)
        worksheet = workbook.add_worksheet()

        row = 0
        col = 0

        for words_list in table[0].values():
            for i in range(len(words_list)):
                worksheet.write(row, col, words_list[i])
                col += 1
            row += 1
            col = 0

        workbook.close()
        return True

    except Exception as error:
        log_error('saving_to_xlsx', file_path_to_save, error)
        return f'Не удалось сохранить файл {error}'

